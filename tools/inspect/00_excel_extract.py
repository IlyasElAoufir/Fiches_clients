# -*- coding: utf-8 -*-
"""
PHASE 1 — Extraction normalisee et EXPURGEE de `Fiches Clients.xlsx`.

Principes :
  - lecture seule du classeur ;
  - ALLOWLIST de libelles : seuls les champs explicitement autorises sont
    conserves. Tout le reste est ignore (pas de blacklist fragile) ;
  - filet de securite supplementaire : DENY par motif (mdp, password, md5,
    key, secret, apikey...) applique APRES l'allowlist ;
  - la feuille "Mdp Admin" n'est jamais ouverte.

Sortie : data/excel-clients.json
"""
import os
import sys
import io
import json
import re
import unicodedata
from pathlib import Path

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")
from openpyxl import load_workbook  # noqa: E402

# Chemin du classeur. Jamais en dur dans le depot : il est lu depuis la
# variable FICHES_CLIENTS_XLSX (environnement ou .env.local, non versionne).
def _xlsx_path():
    value = os.environ.get("FICHES_CLIENTS_XLSX", "")
    if not value:
        env_file = Path(__file__).resolve().parents[2] / ".env.local"
        if env_file.exists():
            for line in env_file.read_text(encoding="utf-8").splitlines():
                line = line.strip()
                if line.startswith("FICHES_CLIENTS_XLSX="):
                    value = line.split("=", 1)[1].strip()
                    break
    if not value:
        sys.exit(
            "ERREUR: chemin du classeur inconnu.\n"
            "Renseignez FICHES_CLIENTS_XLSX dans .env.local, par exemple :\n"
            "  FICHES_CLIENTS_XLSX=C:\\dossier\\Fiches Clients.xlsx"
        )
    return Path(value)


XLSX = _xlsx_path()
OUT = Path(__file__).resolve().parents[2] / "data" / "excel-clients.json"

SKIP_SHEETS = {"Mdp Admin", "Index", "Mod\u00e8le fiche (nouveau)"}

# ---------------------------------------------------------------------------
# Filet de securite : rien contenant ces motifs ne sort jamais
# ---------------------------------------------------------------------------
DENY = re.compile(
    r"(mdp|mot\s*de\s*passe|password|passwd|\bpwd\b|md5|hash|"
    r"share\s*key|\bkey\b|cl\u00e9\s*api|cle\s*api|apikey|api\s*key|"
    r"secret|token|certificat|questions?\s*de\s*s\u00e9curit\u00e9|auth\s*2)",
    re.IGNORECASE,
)


def strip_accents(s: str) -> str:
    return "".join(c for c in unicodedata.normalize("NFD", s)
                   if unicodedata.category(c) != "Mn")


def key(s: str) -> str:
    """Normalise un libelle pour comparaison : minuscules, sans accent/ponctuation."""
    s = strip_accents(str(s)).lower()
    s = s.replace("*", " ")
    s = re.sub(r"[^a-z0-9]+", " ", s)
    return re.sub(r"\s+", " ", s).strip()


# ---------------------------------------------------------------------------
# ALLOWLIST : libelle normalise -> (section, champ canonique, a_2_env)
#   a_2_env=True  -> la colonne D = PROD, la colonne E = INT
#   a_2_env=False -> seule la colonne D est lue
# ---------------------------------------------------------------------------
ALLOW = {
    # --- Identite -----------------------------------------------------------
    "nom organisme":              ("identite", "nomOrganisme", False),
    "site web":                   ("identite", "siteWeb", False),
    "adresse":                    ("identite", "adresse", False),
    "code postal":                ("identite", "codePostal", False),
    "ville":                      ("identite", "ville", False),
    "tel contact organisame":     ("identite", "telOrganisme", False),
    "tel contact organisme":      ("identite", "telOrganisme", False),
    "email contact organisme":    ("identite", "emailOrganisme", False),
    "informations organisme":     ("identite", "lienInfosOrganisme", False),

    # --- Application --------------------------------------------------------
    "prefix trigramme client":    ("application", "codeClient", False),
    "prefix user sql":            ("application", "codeClient", False),
    "id organisation":            ("application", "idOrganisation", False),
    "erp du client":              ("application", "erp", False),
    "utilisation novamap":        ("application", "utilisationNovamap", False),
    "type connexion utilisateurs": ("application", "typeConnexion", False),
    "version connecteur":         ("application", "versionConnecteur", True),
    "utilisateur novamap de test": ("application", "utilisateurTest", False),
    "referent novamap":           ("application", "referentNovamap", False),
    "referent metier":            ("application", "referentMetier", False),
    "chef de projet":             ("application", "chefDeProjet", False),

    # --- Environnements -----------------------------------------------------
    "nom base de donnees":        ("environnements", "baseNom", False),
    "base de donnees":            ("environnements", "base", True),
    "store":                      ("environnements", "store", True),
    "azure share name":           ("environnements", "azureShareName", True),
    "vm dediee":                  ("environnements", "vmDediee", False),
    "url conf":                   ("environnements", "urlConf", False),
    # NB: les libelles "Base", "Serveur" et "Host" ne sont volontairement PAS
    # dans l'allowlist : ils n'apparaissent que dans les sections d'acces aux
    # vues SQL, ou ils sont peu fiables (valeurs PROD recopiees en INT).

    # --- Materiel -----------------------------------------------------------
    "tablette":                   ("materiel", "typeTablette", False),
    "version os tablette":        ("materiel", "versionOsTablette", False),
    "modeles ipad":               ("materiel", "modelesIpad", False),
    "plateforme":                 ("materiel", "plateforme", False),

    # --- Divers -------------------------------------------------------------
    "notes":                      ("notes", "notes", False),
    "memo client":                ("notes", "memoClient", False),
}

# Libelles de contact, traites a part (rattaches au role qui precede)
CONTACT_FIELDS = {
    "nom": "nom",
    "prenom": "prenom",
    "poste": "poste",
    "mail": "email",
    "email": "email",
    "mail exploitation": "emailExploitation",
    "tel": "tel",
    "n telephone": "tel",
    "n portable": "portable",
}

# Sections (colonne B) qui ouvrent un bloc de contacts
CONTACT_SECTIONS = re.compile(r"contact", re.IGNORECASE)

# Sections entierement ignorees : credentials d'acces a des outils tiers,
# ou donnees non exploitables par l'application.
SKIP_SECTIONS = re.compile(
    r"(VUES?|VUE EXCEL|NOVADOC|MANTIS|APPLE|SMS|PROCSTOCK|MERGE|"
    r"ACCES ECRITURE|CERTIFICAT|CONFIGURATION MAIL)",
    re.IGNORECASE,
)

MODULES = ["NKeep", "NBuild", "NSupply (Old)", "Nsupply 2", "NConnect",
           "Recensement Eqt.", "EDL", "R\u00e9clamations", "DI",
           "Constat Complt.", "Pr\u00e9-visites", "Controles S\u00e9curit\u00e9",
           "Controles Securit\u00e9", "Controles Qualit\u00e9s"]

# Valeurs placeholder du modele : a considerer comme "non renseigne"
PLACEHOLDERS = {"oui/non", "android/ios", "xxx", "n/a", "nsp", "a definir",
                "\u00e0 d\u00e9finir", "inconnu", "-int", "aucun", "",
                "nom client"}


def clean(v):
    if v is None:
        return ""
    s = re.sub(r"\s+", " ", str(v)).strip()
    return "" if s.lower() in PLACEHOLDERS else s


def truthy_module(v: str):
    lv = v.lower()
    if lv in ("oui", "yes", "1", "x"):
        return True
    if lv in ("non", "no", "0"):
        return False
    return None


# ---------------------------------------------------------------------------
def parse_sheet(ws, name):
    out = {
        "sheet": name,
        "clientName": re.sub(r"\s*-\s*A SUPPRIMER\s*$", "", name).strip(),
        "aSupprimer": "A SUPPRIMER" in name.upper(),
        "format": None,
        "identite": {},
        "application": {},
        "environnements": {},
        "materiel": {},
        "notes": {},
        "contacts": [],
        "modules": {},
        "_rejetes": 0,
    }

    rows = list(ws.iter_rows(values_only=True))
    current_section = ""       # colonne B
    current_role = ""          # dernier libelle "role" rencontre
    current_contact = None
    env_context = ""           # "prod" / "int" / "" selon le sous-titre courant
    is_nouveau = False

    def flush_contact():
        nonlocal current_contact
        if current_contact and any(current_contact.get(k) for k in
                                   ("nom", "prenom", "email", "tel", "portable")):
            out["contacts"].append(current_contact)
        current_contact = None

    for row in rows:
        b = clean(row[1]) if len(row) > 1 else ""
        c = clean(row[2]) if len(row) > 2 else ""
        d = clean(row[3]) if len(row) > 3 else ""
        e = clean(row[4]) if len(row) > 4 else ""

        if b:
            flush_contact()
            current_section = b
            current_role = ""
            env_context = ""

        if not c:
            continue

        # Sections d'acces a des outils tiers : ignorees en bloc
        if current_section and SKIP_SECTIONS.search(current_section):
            continue

        kc = key(c)
        if kc == "prefix trigramme client":
            is_nouveau = True

        # --- Modules --------------------------------------------------------
        if c in MODULES:
            t = truthy_module(d)
            if t is not None:
                out["modules"][c] = t
            continue

        # --- Contexte d'environnement (sous-titres) -------------------------
        if kc in ("production", "environnement production"):
            env_context = "prod"
            continue
        if kc in ("integration", "recette", "environnement integration",
                  "environnement recette"):
            env_context = "int"
            continue
        if kc in ("formation", "environnement formation"):
            env_context = "form"
            continue

        # --- Champs de contact ---------------------------------------------
        if kc in CONTACT_FIELDS and CONTACT_SECTIONS.search(current_section or ""):
            if current_contact is None:
                current_contact = {"role": current_role or "Contact"}
            field = CONTACT_FIELDS[kc]
            if d and not DENY.search(c):
                current_contact[field] = d
            continue

        # --- Un libelle non reconnu dans une section contact = un role ------
        if CONTACT_SECTIONS.search(current_section or "") and not d and not e:
            flush_contact()
            current_role = c
            continue

        # --- ALLOWLIST ------------------------------------------------------
        if kc in ALLOW:
            section, field, two_env = ALLOW[kc]

            # Filet de securite : le libelle ne doit rien evoquer de sensible
            if DENY.search(c):
                out["_rejetes"] += 1
                continue

            if two_env:
                slot = out[section].setdefault(field, {})
                if e:
                    # Format NOUVEAU : colonne D = Production, colonne E = Recette
                    if d:
                        slot["prod"] = d
                    slot["int"] = e
                elif d:
                    # Format LEGACY : une seule colonne, le sous-titre courant
                    # ("Environnement Production" / "... Recette") donne l'env.
                    slot[env_context or "prod"] = d
            else:
                if d:
                    out[section][field] = d
                elif e:
                    out[section][field] = e
            continue

        out["_rejetes"] += 1

    flush_contact()

    out["format"] = "NOUVEAU" if is_nouveau else "LEGACY"

    # Validation du code client : trigramme court, sinon la cellule Excel
    # contient autre chose (ex. un login de vue SQL "LF-SQL-VIEW").
    code = out["application"].get("codeClient", "")
    if code and (len(code) > 6 or not re.fullmatch(r"[A-Za-z0-9]+", code)):
        out.setdefault("_anomalies", []).append(
            f"codeClient rejete (valeur invalide dans l'Excel) : {code!r}")
        out["application"].pop("codeClient")

    # Base PROD canonique : "Base de donnees" (prod) sinon "Nom base de donnees"
    env = out["environnements"]
    prod = env.get("base", {}).get("prod") or env.get("baseNom") or ""
    integ = env.get("base", {}).get("int") or ""
    # Si seule l'INT est connue, la PROD se deduit par convention
    if not prod and integ.endswith("-int"):
        prod = integ[: -len("-int")]
        env["baseDeduite"] = True
    if not integ and prod and prod != "NovamapSharedDB":
        integ = prod + "-int"
        env["baseIntDeduite"] = True
    out["resolved"] = {"dbProd": prod, "dbInt": integ}
    return out


def main():
    wb = load_workbook(XLSX, data_only=True, read_only=True)
    clients = []
    for name in wb.sheetnames:
        if name in SKIP_SHEETS:
            continue
        clients.append(parse_sheet(wb[name], name))
    wb.close()

    OUT.parent.mkdir(parents=True, exist_ok=True)
    OUT.write_text(json.dumps(clients, ensure_ascii=False, indent=2),
                   encoding="utf-8")

    # --- Controle anti-fuite sur le JSON produit ---------------------------
    raw = OUT.read_text(encoding="utf-8")
    leaks = [m.group(0) for m in DENY.finditer(raw)]
    print(f"{len(clients)} fiches extraites -> {OUT}")
    print(f"Controle anti-fuite : {len(leaks)} occurrence(s) de motif sensible "
          f"{'!! A VERIFIER: ' + str(set(leaks)) if leaks else 'OK'}")

    print()
    print(f"{'CLIENT':<32} {'FMT':<8} {'CODE':<6} {'ORGID':<6} "
          f"{'DB PROD':<17} {'DB INT':<19} {'CT':<3} {'DED'}")
    print("-" * 110)
    for c in clients:
        env = c["environnements"]
        ded = []
        if env.get("baseDeduite"):
            ded.append("prod")
        if env.get("baseIntDeduite"):
            ded.append("int")
        print(f"{c['clientName'][:31]:<32} {c['format']:<8} "
              f"{c['application'].get('codeClient', '')[:5]:<6} "
              f"{str(c['application'].get('idOrganisation', ''))[:5]:<6} "
              f"{c['resolved']['dbProd'][:16]:<17} "
              f"{c['resolved']['dbInt'][:18]:<19} "
              f"{len(c['contacts']):<3} {','.join(ded)}")


if __name__ == "__main__":
    main()
